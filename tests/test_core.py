import importlib.util
import json
import pathlib
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

SPEC = importlib.util.spec_from_file_location(
    "report_core",
    pathlib.Path(__file__).resolve().parents[1] / "scripts/report_core.py",
)
core = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(core)


class CoreRegression(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = pathlib.Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def excel(self, rows, filename="synthetic.xlsx", **kwargs):
        path = self.directory / filename
        pd.DataFrame(rows).to_excel(path, index=False, **kwargs)
        return path

    def row(
        self,
        waybill="001234",
        ship="2026-09-01 09:00",
        sign="2026-09-02 09:00",
        **kwargs,
    ):
        row = {
            "运单号": waybill,
            "寄件时间": ship,
            "签收时间": sign,
            "目的省份": "测试省",
            "件数": 2,
            "实际重量": 4.5,
            "结算重量": 5.0,
            "体积": 0.2,
            "备注": None,
        }
        row.update(kwargs)
        return row

    def analyze(self, rows, as_of="2026-09-10 09:00", config=None):
        return core.analyze(
            core.load_many([self.excel(rows)]), as_of=as_of, config=config
        )

    def test_blank_text_is_not_exception_or_province(self):
        result = self.analyze([self.row(目的省份=None)])
        self.assertEqual(result["异常单数"], 0)
        self.assertEqual(result["review_count"], 0)
        self.assertEqual(result["省份数"], 0)
        self.assertEqual(result["运单数"], 1)

    def test_ids_keep_text_leading_zeros_and_cell_padding(self):
        path = self.excel([self.row(), self.row(waybill="2"), self.row(waybill="NA")])
        book = load_workbook(path)
        book.active["A3"] = 123
        book.active["A3"].number_format = "000000"
        book.save(path)
        frame = core.load_many([path])
        self.assertEqual(frame["运单号"].tolist(), ["001234", "000123", "NA"])

    def test_required_columns_and_rows_fail_closed(self):
        with self.assertRaisesRegex(ValueError, "缺少必需列"):
            core.load_many([self.excel([{"寄件时间": "2026-09-01"}])])
        with self.assertRaisesRegex(ValueError, "没有有效运单"):
            core.load_many([self.excel([self.row(ship="not a date")])])

    def test_duplicate_removed_and_conflict_rejected(self):
        first = self.excel([self.row()], "first.xlsx")
        frame = core.load_many([first, first])
        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.attrs["quality"]["duplicate_rows"], 1)
        self.assertEqual(frame.attrs["quality"]["input_rows"], 2)
        self.assertEqual(
            sum(item["valid_rows"] for item in frame.attrs["quality"]["row_counts"]), 1
        )
        self.assertEqual(
            frame.attrs["quality"]["input_files"][0]["columns"]["运单号"], "运单号"
        )
        duplicate_issues = [
            issue
            for issue in frame.attrs["quality"]["issues"]
            if issue["reason"] == "重复运单，已去除"
        ]
        self.assertEqual(len(duplicate_issues), 1)
        self.assertEqual(duplicate_issues[0]["field"], "运单号")
        self.assertEqual(duplicate_issues[0]["value"], "001234")
        second = self.excel([self.row(件数=3)], "second.xlsx")
        with self.assertRaisesRegex(ValueError, "001234.*件数") as caught:
            core.load_many([first, second])
        self.assertEqual(len(caught.exception.quality["conflicts"]), 1)

    def test_mature_cohort_counts_overdue_unsigned(self):
        rows = [self.row("S001")] + [self.row(f"U{i:03}", sign=None) for i in range(99)]
        result = self.analyze(rows)
        self.assertEqual(result["签收率"], 1.0)
        self.assertEqual(result["72h率"], 100.0)
        self.assertEqual(result["sla_rate"], 1.0)
        self.assertEqual(result["sla_eligible"], 100)
        self.assertEqual(result["sla_overdue_unsigned"], 99)
        self.assertEqual(result["province_sla"]["测试省"]["rate"], 1.0)

    def test_future_sign_cutoff_and_future_ship(self):
        result = self.analyze(
            [
                self.row("S001", sign="2026-09-05 09:00"),
                self.row("S002", ship="2026-09-08 09:00", sign=None),
            ],
            as_of="2026-09-04 09:00",
        )
        self.assertEqual(result["运单数"], 1)
        self.assertEqual(result["已签收单数"], 0)
        self.assertEqual(result["sla_overdue_unsigned"], 1)
        self.assertEqual(result["quality"]["future_sign_rows"], 1)
        self.assertEqual(result["quality"]["as_of_excluded_rows"], 1)

    def test_pending_cohort_no_success_credit(self):
        result = self.analyze([self.row(sign=None)], as_of="2026-09-02 09:00")
        self.assertEqual(result["sla_pending"], 1)
        self.assertIsNone(result["sla_rate"])
        self.assertIsNone(result["72h率"])
        self.assertIsNone(result["平均时效"])
        self.assertIsNone(result["province_sla"]["测试省"]["rate"])
        for obsolete in ("期末日", "prov_recent_unship", "期末截点在途"):
            self.assertNotIn(obsolete, result)

    def test_date_precision_excludes_exact_hour_metrics(self):
        result = self.analyze(
            [
                self.row("DATE_SHIP", ship="2026-09-01", sign="2026-09-02 14:00"),
                self.row("DATE_SIGN", sign="2026-09-02"),
            ]
        )
        self.assertEqual(result["已签收单数"], 2)
        self.assertEqual(result["sla_precision_excluded"], 2)
        self.assertEqual(result["sla_eligible"], 0)
        self.assertIsNone(result["72h率"])
        self.assertIsNone(result["平均时效"])

    def test_excel_date_format_retains_precision(self):
        path = self.excel([self.row(ship=pd.Timestamp("2026-09-01"))])
        book = load_workbook(path)
        book.active["B2"].number_format = "yyyy-mm-dd"
        book.save(path)
        frame = core.load_many([path])
        self.assertEqual(frame.iloc[0]["__ship_precision"], "date")
        self.assertIsNone(core.analyze(frame, as_of="2026-09-10")["sla_rate"])

    def test_zero_duration_and_negative_quarantine(self):
        result = self.analyze(
            [
                self.row("ZERO", sign="2026-09-01 09:00"),
                self.row("NEG", sign="2026-09-01 08:00"),
            ]
        )
        self.assertEqual(result["已签收单数"], 1)
        self.assertEqual(result["有效时效单数"], 1)
        self.assertEqual(result["时效分布"]["24h内"], 1)
        self.assertEqual(result["24h率"], 100)
        self.assertEqual(result["平均时效"], 0)
        self.assertEqual(result["sla_invalid_excluded"], 1)
        self.assertEqual(result["签收待核实单数"], 1)
        self.assertEqual(result["可确认未签收单数"], 0)
        self.assertEqual(result["时效隔离票数"], 1)
        self.assertTrue(
            any(
                issue["field"] == "签收时间" and "早于寄件" in issue["reason"]
                for issue in result["quality"]["issues"]
            )
        )

    def test_invalid_numeric_stays_missing_and_piece_assumptions_survive_merge(self):
        first = self.excel([self.row("A", 实际重量="bad", 件数="bad")], "a.xlsx")
        second_row = self.row("B")
        second_row.pop("件数")
        second = self.excel([second_row], "b.xlsx")
        frame = core.load_many([first, second])
        result = core.analyze(frame, as_of="2026-09-10")
        self.assertTrue(result["件数按单计"])
        self.assertEqual(result["推定件数票数"], 1)
        self.assertTrue(pd.isna(frame.iloc[0]["实际重量"]))
        self.assertTrue(pd.isna(frame.iloc[0]["件数"]))
        self.assertEqual(result["总件数"], 1)

    def test_exception_priority_negation_and_review(self):
        cases = {
            "在途丢失": core.OTHER_LABEL,
            "中转中，延误3天": "目的分拨中转延迟",
            "正常签收": None,
            "分拨正常，无异常": None,
            "无需预约": None,
            "分拨无延误": None,
            "未发生丢失": None,
            "运输中，没有破损": core.IN_TRANSIT_LABEL,
            "派送途中": core.IN_TRANSIT_LABEL,
            "待了解": core.REVIEW_LABEL,
            "周末不收货": "客户预约延迟",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(core.classify_exception(text), expected)
        result = self.analyze(
            [self.row("A", 备注="待了解"), self.row("B", 备注="周末不收货")]
        )
        self.assertEqual(result["review_count"], 1)
        self.assertEqual(result["异常单数"], 1)
        self.assertEqual(result["异常_客户约定"], 1)

    def test_province_sla_override_and_timezone(self):
        result = self.analyze(
            [self.row(sign="2026-09-03 09:00")],
            config={"province_sla_hours": {"测试省": 24}},
        )
        self.assertEqual(result["sla_rate"], 0)
        self.assertEqual(result["province_sla"]["测试省"]["sla_hours"], 24)
        self.assertEqual(str(result["as_of"].tz), "Asia/Shanghai")

    def test_logistics_exception_excludes_customer_agreement(self):
        result = self.analyze(
            [
                self.row("A", 备注="周末不收货"),
                self.row("B", 备注="中转延误3天"),
                self.row("C", 备注="正常签收"),
            ]
        )
        self.assertEqual(result["异常单数"], 2)
        self.assertEqual(result["客户约定单数"], 1)
        self.assertEqual(result["物流异常单数"], 1)
        self.assertAlmostEqual(result["物流异常率"], 100 / 3)
        self.assertAlmostEqual(result["异常率"], 200 / 3)

    def test_invalid_sign_evidence_is_disclosed_separately(self):
        result = self.analyze(
            [self.row("INVALID", sign="invalid"), self.row("UNSIGNED", sign=None)]
        )
        self.assertEqual(result["已签收单数"], 0)
        self.assertEqual(result["未签收"], 2)
        self.assertEqual(result["可确认未签收单数"], 1)
        self.assertEqual(result["签收待核实单数"], 1)
        self.assertEqual(result["sla_invalid_excluded"], 1)
        self.assertEqual(result["sla_eligible"], 1)
        self.assertEqual(result["sla_overdue_unsigned"], 1)

    def test_directory_lock_files_and_existing_comma_filename(self):
        path = self.excel([self.row()], "a,b.xlsx")
        self.excel([self.row()], "~$temporary.xlsx")
        self.assertEqual(core.resolve_files(str(path)), [str(path)])
        self.assertEqual(core.resolve_files(str(self.directory)), [str(path)])
        self.assertEqual(
            core.resolve_files([str(path), str(path)]), [str(path), str(path)]
        )

    def test_configuration_validation(self):
        for config in (
            {"sla_hours": 0},
            {"sla_hours": True},
            {"sla_target": 101},
            {"timezone": "Unknown/Place"},
            {"column_mapping": {"unknown": "x"}},
            {"province_sla_hours": {"测试省": -1}},
            {"mystery": True},
        ):
            with self.subTest(config=config), self.assertRaises(ValueError):
                core.validate_config(config)
        path = self.directory / "config.json"
        path.write_text(
            json.dumps({"client": "测试", "response_hours": 2}), encoding="utf-8"
        )
        config = core.load_config(path)
        self.assertEqual(config["client"], "测试")
        self.assertEqual(config["sla_hours"], 72)

    def test_all_future_shipments_empty_cohort(self):
        result = self.analyze([self.row()], as_of="2026-08-01 09:00")
        self.assertEqual(result["运单数"], 0)
        self.assertIsNone(result["签收率"])
        self.assertIsNone(result["sla_rate"])
        self.assertTrue(result["daily"].empty)

    def test_named_sheet_header_and_column_mapping(self):
        path = self.excel(
            [
                {
                    " ID ": "0001",
                    " Ship ": "2026-09-01 09:00",
                    " Sign ": "2026-09-02 09:00",
                }
            ],
            sheet_name="运单",
            startrow=2,
        )
        frame = core.load_many(
            [path],
            sheet="运单",
            header=2,
            column_mapping={"运单号": "ID", "寄件时间": "Ship", "签收时间": "Sign"},
        )
        self.assertEqual(frame.iloc[0]["运单号"], "0001")
        self.assertEqual(frame.iloc[0]["__source_row"], 4)
        self.assertTrue(frame.iloc[0]["__pieces_assumed"])
        self.assertEqual(
            frame.attrs["quality"]["input_files"][0]["columns"]["运单号"], "ID"
        )

    def test_three_duplicates_keep_explicit_piece_evidence(self):
        paths = [
            self.excel([self.row(件数=None)], "missing_first.xlsx"),
            self.excel([self.row(件数=1)], "explicit.xlsx"),
            self.excel([self.row(件数=None)], "missing_last.xlsx"),
        ]
        frame = core.load_many(paths)
        self.assertEqual(len(frame), 1)
        self.assertFalse(frame.iloc[0]["__pieces_assumed"])
        self.assertEqual(frame.attrs["quality"]["duplicate_rows"], 2)
        self.assertEqual(
            sum(item["valid_rows"] for item in frame.attrs["quality"]["row_counts"]), 1
        )
        self.assertEqual(core.analyze(frame, as_of="2026-09-10")["推定件数票数"], 0)

    def test_date_only_cross_day_reverse_is_invalid_but_same_day_is_uncertain(self):
        frame = core.load_many(
            [
                self.excel(
                    [
                        self.row("REVERSE", ship="2026-09-21 09:00", sign="2026-09-20"),
                        self.row(
                            "SAME_DAY", ship="2026-09-21 09:00", sign="2026-09-21"
                        ),
                    ]
                )
            ]
        )
        self.assertFalse(frame.iloc[0]["__time_valid"])
        self.assertTrue(frame.iloc[1]["__time_valid"])
        result = core.analyze(frame, as_of="2026-09-30")
        self.assertEqual(result["已签收单数"], 1)
        self.assertEqual(result["签收待核实单数"], 1)
        self.assertEqual(result["sla_invalid_excluded"], 1)
        self.assertEqual(result["sla_precision_excluded"], 1)
        self.assertEqual(
            result["_data"]["SLA样本状态"].tolist(),
            ["time_invalid", "precision_insufficient"],
        )
        # analyze also validates manually supplied/corrected normalized records.
        frame["__time_valid"] = True
        rerun = core.analyze(frame, as_of="2026-09-30")
        self.assertEqual(rerun["sla_invalid_excluded"], 1)
        self.assertEqual(rerun["已签收单数"], 1)

    def test_row_sla_states_reconcile_cohort_counts(self):
        rows = [
            self.row("ONTIME"),
            self.row("LATE", sign="2026-09-05 09:00"),
            self.row("UNSIGNED", sign=None),
            self.row("PENDING", ship="2026-09-09 09:00", sign=None),
            self.row("DATE", ship="2026-09-01"),
            self.row("INVALID", sign="invalid"),
        ]
        result = self.analyze(rows)
        states = result["_data"].set_index("运单号")["SLA样本状态"].to_dict()
        self.assertEqual(
            states,
            {
                "ONTIME": "on_time",
                "LATE": "late_signed",
                "UNSIGNED": "overdue_unsigned",
                "PENDING": "pending",
                "DATE": "precision_insufficient",
                "INVALID": "time_invalid",
            },
        )
        counts = result["_data"]["SLA样本状态"].value_counts()
        self.assertEqual(
            result["sla_eligible"],
            counts["on_time"] + counts["late_signed"] + counts["overdue_unsigned"],
        )
        self.assertEqual(result["sla_ontime"], counts["on_time"])
        self.assertTrue(
            pd.isna(result["_data"].set_index("运单号").loc["DATE", "SLA截止时间"])
        )

    def test_exception_status_groups_cover_every_exception_record(self):
        remarks = [
            "周末不收货",
            "未赶上清仓",
            "中转延误",
            "时效顺延",
            "缺送货单",
            "需预约",
            "网点异常",
            "在途丢失",
        ]
        result = self.analyze(
            [self.row(f"E{i}", 备注=remark) for i, remark in enumerate(remarks)]
        )
        self.assertEqual(
            sum(
                result[key]
                for key in (
                    "异常_客户约定",
                    "异常_物流环节",
                    "异常_单证操作",
                    "异常_其他",
                )
            ),
            result["异常单数"],
        )
        self.assertEqual(
            result["客户约定单数"] + result["物流异常单数"], result["异常单数"]
        )

    def test_missing_sign_or_remark_columns_produces_unknown_rates(self):
        row = self.row()
        row.pop("签收时间")
        row.pop("备注")
        frame = core.load_many([self.excel([row])])
        self.assertFalse(frame.iloc[0]["__sign_available"])
        self.assertFalse(frame.iloc[0]["__remark_available"])
        self.assertFalse(frame.iloc[0]["__time_valid"])
        result = core.analyze(frame, as_of="2026-09-10")
        self.assertEqual(result["签收字段覆盖票数"], 0)
        self.assertEqual(result["备注字段覆盖票数"], 0)
        self.assertIsNone(result["签收率"])
        self.assertIsNone(result["异常率"])
        self.assertIsNone(result["物流异常率"])
        self.assertIsNone(result["sla_rate"])
        self.assertEqual(result["sla_eligible"], 0)
        self.assertEqual(result["sla_overdue_unsigned"], 0)
        self.assertEqual(result["异常单数"], 0)
        self.assertTrue(pd.isna(result["exc_daily"].iloc[0]["异常率"]))
        self.assertTrue(
            any(
                issue["field"] == "签收时间" and "未提供" in issue["reason"]
                for issue in result["quality"]["issues"]
            )
        )

    def test_mixed_field_coverage_preserves_counts_but_not_total_rates(self):
        covered = self.excel(
            [self.row("COVERED", sign=None, 备注="中转延误")], "covered.xlsx"
        )
        missing_row = self.row("MISSING")
        missing_row.pop("签收时间")
        missing_row.pop("备注")
        missing = self.excel([missing_row], "missing.xlsx")
        result = core.analyze(core.load_many([covered, missing]), as_of="2026-09-10")
        self.assertEqual(result["运单数"], 2)
        self.assertEqual(result["签收字段覆盖票数"], 1)
        self.assertEqual(result["备注字段覆盖票数"], 1)
        self.assertIsNone(result["签收率"])
        self.assertIsNone(result["异常率"])
        self.assertIsNone(result["物流异常率"])
        self.assertEqual(result["异常单数"], 1)
        self.assertEqual(result["sla_eligible"], 1)
        self.assertEqual(result["sla_overdue_unsigned"], 1)
        self.assertEqual(result["sla_invalid_excluded"], 1)
        self.assertEqual(result["可确认未签收单数"], 1)
        self.assertEqual(result["签收待核实单数"], 1)

    def test_duplicate_field_availability_conflicts(self):
        missing_row = self.row(sign=None)
        missing_row.pop("签收时间")
        missing = self.excel([missing_row], "missing.xlsx")
        covered = self.excel([self.row(sign=None)], "covered.xlsx")
        with self.assertRaisesRegex(ValueError, "__sign_available"):
            core.load_many([missing, covered])
        missing_remark_row = self.row()
        missing_remark_row.pop("备注")
        missing_remark = self.excel([missing_remark_row], "missing_remark.xlsx")
        with self.assertRaisesRegex(ValueError, "__remark_available"):
            core.load_many(
                [missing_remark, self.excel([self.row()], "covered_remark.xlsx")]
            )

    def test_manual_frame_defaults_to_available_fields(self):
        frame = core.load_many([self.excel([self.row(sign=None)])]).drop(
            columns=["__sign_available", "__remark_available"]
        )
        result = core.analyze(frame, as_of="2026-09-10")
        self.assertEqual(result["签收字段覆盖票数"], 1)
        self.assertEqual(result["备注字段覆盖票数"], 1)
        self.assertEqual(result["签收率"], 0)
        self.assertEqual(result["物流异常率"], 0)
        self.assertEqual(result["sla_overdue_unsigned"], 1)


if __name__ == "__main__":
    unittest.main()
